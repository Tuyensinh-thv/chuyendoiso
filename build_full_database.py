import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json

def create_hvu_database():
    with open(r"C:\Users\Admin\.gemini\antigravity-ide\brain\b61d93d4-185a-4b60-a339-42e883c35131\scratch\official_tasks_doc.json", "r", encoding="utf-8") as f:
        tasks = json.load(f)

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    navy_fill = PatternFill(start_color="0B2545", end_color="0B2545", fill_type="solid")
    red_fill = PatternFill(start_color="BE1E2D", end_color="BE1E2D", fill_type="solid")
    subheader_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    light_gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    accent_gold_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Arial", size=10, color="1E293B")
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # 1. BẢNG 01: NHIỆM VỤ NQ57 (57 NHIỆM VỤ)
    ws1 = wb.create_sheet("01_Nhiem_Vu_NQ57")
    headers1 = [
        "Ma_NV", "Nhom_Ke_Hoach", "Tieu_de_Nhiem_vu", "Don_vi_Chu_tri", 
        "Don_vi_Phoi_hop", "San_pham_Dau_ra", "Thoi_han_Van_Ban", "Thoi_han_Chuan", 
        "Trang_thai", "Muc_do_Uu_tien", "Tien_do_%", "Nguoi_phu_trach", 
        "Link_Minh_Chung", "Ghi_chu_Noi_bo", "Y_kien_Chi_dao", "Ngay_Cap_nhat"
    ]
    ws1.append(headers1)

    group_code_map = {
        "Kế hoạch 197 (100 ngày xử lý điểm nghẽn CĐS)": "KH197",
        "Kế hoạch 200 (Chiến dịch 90 ngày Kho dữ liệu số)": "KH200",
        "Kế hoạch 154 & 201 (Bình dân học vụ số)": "BDHVS",
        "Kế hoạch 172 (Triển khai NQ 113 Đảng ủy)": "NQ113",
        "Kế hoạch 188 (KHCN, ĐMST & CĐS 2026-2027)": "KH188",
        "QĐ 1034 (Hội thảo, tọa đàm khoa học CĐS)": "QD1034",
        "Nhiệm vụ UBND Tỉnh giao (Hoàn thành T12/2026)": "UBND_TINH"
    }

    for t in tasks:
        ws1.append([
            t["id"],
            t["nhomKeHoach"],
            t["tenNhiemVu"],
            t["donViChuTri"],
            t["donViPhoiHop"],
            t["sanPhamDauRa"],
            t["thoiHanVanBan"],
            t["thoiHan"],
            "Chưa thực hiện",
            t["mucDoUuTien"],
            0,
            t["nguoiPhuTrach"],
            "",
            t["ghiChuNoiBo"],
            "",
            t["ngayCapNhat"]
        ])

    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws1.iter_rows(min_row=2, max_row=len(tasks)+1, min_col=1, max_col=len(headers1)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            if cell.column_letter in ['A', 'G', 'H', 'I', 'J', 'K', 'P']:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 2. BẢNG 02: NHÓM KẾ HOẠCH
    ws2 = wb.create_sheet("02_Nhom_Ke_Hoach")
    headers2 = ["Ma_Nhom_KH", "Ten_Ke_Hoach", "So_Quyet_Dinh", "Ngay_Ban_Hanh", "Co_Quan_Ban_Hanh", "So_Nhiem_Vu", "Ghi_Chu"]
    ws2.append(headers2)
    plan_groups_data = [
        ["KH197", "Kế hoạch 197 (100 ngày xử lý điểm nghẽn CĐS)", "197/KH-ĐHHV", "2026-06-15", "Hiệu trưởng ĐHHV", 11, "Hành động 100 ngày xử lý các điểm nghẽn về chuyển đổi số"],
        ["KH200", "Kế hoạch 200 (Chiến dịch 90 ngày Kho dữ liệu số)", "200/KH-ĐHHV", "2026-06-20", "Hiệu trưởng ĐHHV", 21, "Chiến dịch 90 ngày làm sạch, tạo lập và đưa vào sử dụng Kho dữ liệu số"],
        ["BDHVS", "Kế hoạch 154 & 201 (Bình dân học vụ số)", "154/KH-ĐHHV & 201/KH-ĐHHV", "2026-05-10", "Ban Giám hiệu / Đảng ủy", 10, "Phát động phong trào Bình dân học vụ số, kỹ năng số, VNeID"],
        ["NQ113", "Kế hoạch 172 (Triển khai NQ 113 Đảng ủy)", "172/KH-ĐHHV", "2026-08-05", "Đảng ủy ĐHHV", 2, "Thực hiện Nghị quyết số 113-NQ/ĐU về KHCN, ĐMST, CĐS & HTQT đến 2035"],
        ["KH188", "Kế hoạch 188 (KHCN, ĐMST & CĐS 2026-2027)", "188/KH-ĐHHV", "2026-08-22", "Hiệu trưởng ĐHHV", 5, "Hoạt động KHCN, đổi mới sáng tạo, chuyển đổi số năm học 2026-2027"],
        ["QD1034", "QĐ 1034 (Hội thảo, tọa đàm khoa học CĐS)", "1034/QĐ-ĐHHV", "2026-08-14", "Hiệu trưởng ĐHHV", 4, "Tổ chức hội thảo, tọa đàm khoa học năm học 2026-2027"],
        ["UBND_TINH", "Nhiệm vụ UBND Tỉnh giao (Hoàn thành T12/2026)", "VB 10697 & 11028/UBND-KGVX", "2026-08-10", "UBND Tỉnh Phú Thọ", 4, "Nhiệm vụ trọng tâm UBND Tỉnh giao hoàn thành trong tháng 12/2026"]
    ]
    for row in plan_groups_data:
        ws2.append(row)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws2.iter_rows(min_row=2, max_row=len(plan_groups_data)+1, min_col=1, max_col=len(headers2)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # 3. BẢNG 03: DANH MỤC ĐƠN VỊ
    ws3 = wb.create_sheet("03_Danh_Muc_Don_Vi")
    headers3 = ["Ma_DV", "Ten_Don_Vi", "Loai_Don_Vi", "Lanh_Dao_Phu_Trach", "Email_Lien_He"]
    ws3.append(headers3)
    departments_data = [
        ["DV01", "Ban Giám hiệu", "Lãnh đạo trường", "TS. Đỗ Khắc Thanh (Hiệu trưởng)", "thanhdk@hvu.edu.vn"],
        ["DV02", "Tổ Chuyển đổi số Trường ĐHHV", "Tổ chuyên trách", "TS. Đỗ Khắc Thanh (Tổ trưởng)", "chuyendoiso@hvu.edu.vn"],
        ["DV03", "Văn phòng", "Phòng chức năng", "ThS. Hoàng Thị Thu Hương (Chánh VP)", "vanphong@hvu.edu.vn"],
        ["DV04", "Phòng Đào tạo", "Phòng chức năng", "TS. Phạm Đức Thắng (Trưởng phòng)", "daotao@hvu.edu.vn"],
        ["DV05", "Phòng Kế hoạch - Tài chính", "Phòng chức năng", "ThS. Nguyễn Văn Hải (Trưởng phòng)", "khtc@hvu.edu.vn"],
        ["DV06", "Phòng KHCN & HTQT", "Phòng chức năng", "TS. Nguyễn Danh Hướng (Trưởng phòng)", "khcn_htqt@hvu.edu.vn"],
        ["DV07", "Phòng Quản lý sinh viên và học viên", "Phòng chức năng", "ThS. Tạ Duy Hưng (Trưởng phòng)", "qlsv_hv@hvu.edu.vn"],
        ["DV08", "Phòng Khảo thí và Đảm bảo chất lượng", "Phòng chức năng", "TS. Vũ Thị Lan Hương (Trưởng phòng)", "kt_dbcl@hvu.edu.vn"],
        ["DV09", "Trung tâm Học liệu và Truyền thông", "Đơn vị phục vụ", "ThS. Cao Xuân Chung (Giám đốc)", "hoclieu_tt@hvu.edu.vn"],
        ["DV10", "Trung tâm KN & ĐMST", "Đơn vị trực thuộc", "TS. Trần Thành Vinh (Giám đốc)", "khoinghiep@hvu.edu.vn"],
        ["DV11", "Khoa KTCN", "Khoa chuyên môn", "TS. Trần Đình Chiến (Trưởng Khoa)", "ktcn@hvu.edu.vn"],
        ["DV12", "Khoa Tiếng Trung Quốc", "Khoa chuyên môn", "ThS. Nguyễn Thị Huệ (Trưởng Khoa)", "tiengtrung@hvu.edu.vn"],
        ["DV13", "Khoa NT&TDTT", "Khoa chuyên môn", "TS. Nguyễn Hữu Dũng (Trưởng Khoa)", "nt_tdtt@hvu.edu.vn"],
        ["DV14", "Đoàn Thanh niên", "Đoàn thể", "Đ/c Bí thư Đoàn Trường", "doanthanhnien@hvu.edu.vn"],
        ["DV15", "Hội đồng Thi đua Khen thưởng", "Hội đồng", "Hội đồng TĐKT Nhà trường", "thiduakhenthuong@hvu.edu.vn"],
        ["DV16", "Trạm Y tế", "Đơn vị phục vụ", "Trưởng Trạm Y tế", "tramyte@hvu.edu.vn"],
        ["DV17", "Các chi bộ", "Đảng đoàn", "Bí thư các Chi bộ", "danguy@hvu.edu.vn"],
        ["DV18", "Ban biên tập tạp chí", "Ban chuyên môn", "Tổng biên tập Tạp chí KH&CN", "tapchikhcn@hvu.edu.vn"],
        ["DV19", "Các đơn vị thuộc và trực thuộc", "Toàn trường", "Thủ trưởng các đơn vị", "hvu@hvu.edu.vn"]
    ]
    for row in departments_data:
        ws3.append(row)
    for col_idx in range(1, len(headers3) + 1):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws3.iter_rows(min_row=2, max_row=len(departments_data)+1, min_col=1, max_col=len(headers3)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # 4. BẢNG 04: TÀI KHOẢN NGƯỜI DÙNG
    ws4 = wb.create_sheet("04_Tai_Khoan_Nguoi_Dung")
    headers4 = ["Email", "Ho_Ten", "Don_Vi", "Vai_Tro", "Avatar"]
    ws4.append(headers4)
    users_data = [
        ["thanhdk@hvu.edu.vn", "Đỗ Khắc Thanh", "Ban Giám hiệu", "Lanh_Dao", "👨‍🏫"],
        ["thanhdtp@hvu.edu.vn", "Đặng Thị Phương Thanh", "Phòng Đào tạo", "To_Chuyen_Trach", "👩‍💻"],
        ["sonlh@hvu.edu.vn", "Lê Hồng Sơn", "Phòng Đào tạo / Tổ CNTT", "To_Chuyen_Trach", "👨‍💻"],
        ["kiennt@hvu.edu.vn", "Nguyễn Trung Kiên", "Khoa KT-CN", "Don_Vi", "💻"],
        ["hungdt@hvu.edu.vn", "Đỗ Tất Hưng", "Phòng Đào tạo", "Don_Vi", "🎓"],
        ["linhpd@hvu.edu.vn", "Phùng Duy Linh", "Phòng Khảo thí & ĐBCL", "Don_Vi", "📋"],
        ["gianghx@hvu.edu.vn", "Hoàng Xuân Giang", "Văn phòng", "Don_Vi", "📁"]
    ]
    for row in users_data:
        ws4.append(row)
    for col_idx in range(1, len(headers4) + 1):
        cell = ws4.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws4.iter_rows(min_row=2, max_row=len(users_data)+1, min_col=1, max_col=len(headers4)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto fit column widths
    for sheet in [ws1, ws2, ws3, ws4]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, min(len(val), 50))
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Remove default empty sheet
    wb.remove(default_sheet)

    wb.save(r"d:\xampp\htdocs\chuyendoiso\CSDL_Hoan_Chinh_HVU_NQ57.xlsx")
    print("Saved CSDL_Hoan_Chinh_HVU_NQ57.xlsx successfully with 57 official tasks!")

if __name__ == "__main__":
    create_hvu_database()
